"""Backend tests for InspectAI (auth + inspection pipeline)."""
import os
import time
import io
import pytest
import requests
from pathlib import Path
from openpyxl import load_workbook

BASE_URL = os.environ['REACT_APP_BACKEND_URL'] if 'REACT_APP_BACKEND_URL' in os.environ else None
if not BASE_URL:
    # fall back to reading frontend/.env
    env = Path("/app/frontend/.env").read_text()
    for line in env.splitlines():
        if line.startswith("REACT_APP_BACKEND_URL="):
            BASE_URL = line.split("=", 1)[1].strip()
BASE_URL = BASE_URL.rstrip("/")
API = f"{BASE_URL}/api"

EMAIL = "inspector@indigo.com"
PASSWORD = "inspect123"
SAMPLE_DIR = Path("/app/template/sample_photos")


@pytest.fixture(scope="session")
def token():
    r = requests.post(f"{API}/auth/login", json={"email": EMAIL, "password": PASSWORD}, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    assert "access_token" in data
    assert data["user"]["email"] == EMAIL
    return data["access_token"]


@pytest.fixture(scope="session")
def auth_headers(token):
    return {"Authorization": f"Bearer {token}"}


# ---- Auth ----
class TestAuth:
    def test_login_wrong_password(self):
        r = requests.post(f"{API}/auth/login", json={"email": EMAIL, "password": "wrong"}, timeout=15)
        assert r.status_code == 401

    def test_me_without_token(self):
        r = requests.get(f"{API}/auth/me", timeout=15)
        assert r.status_code == 401

    def test_me_with_token(self, auth_headers):
        r = requests.get(f"{API}/auth/me", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        assert r.json()["email"] == EMAIL


# ---- Inspections pipeline ----
@pytest.fixture(scope="session")
def inspection_id(auth_headers):
    files = []
    for p in sorted(SAMPLE_DIR.iterdir()):
        if p.suffix.lower() in (".png", ".jpg", ".jpeg", ".webp"):
            ct = "image/png" if p.suffix.lower() == ".png" else "image/jpeg"
            files.append(("files", (p.name, p.read_bytes(), ct)))
    assert len(files) >= 10, f"expected sample photos, found {len(files)}"
    r = requests.post(f"{API}/inspections", headers=auth_headers, files=files, timeout=120)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["status"] == "queued"
    assert data["photo_count"] == len(files)
    return data["id"]


class TestInspection:
    def test_upload_returned_id(self, inspection_id):
        assert isinstance(inspection_id, str) and len(inspection_id) > 10

    def test_processing_completes(self, inspection_id, auth_headers):
        statuses = []
        progresses = []
        deadline = time.time() + 240
        final = None
        while time.time() < deadline:
            r = requests.get(f"{API}/inspections/{inspection_id}", headers=auth_headers, timeout=30)
            assert r.status_code == 200
            d = r.json()
            statuses.append(d["status"])
            progresses.append(d.get("progress", 0))
            if d["status"] in ("done", "error"):
                final = d
                break
            time.sleep(5)
        assert final is not None, f"timeout; statuses={statuses}"
        assert final["status"] == "done", f"status={final['status']} err={final.get('error')}"
        assert final["progress"] == 100
        assert "analyzing" in statuses or statuses[0] in ("analyzing", "generating", "done")
        assert len(final.get("products", [])) >= 1
        for prod in final["products"]:
            assert "sheet_name" in prod
            assert "visible_unit_count" in prod
            assert "defects" in prod
            assert "unresolved" in prod
            assert "photos" in prod

    def test_list_inspections(self, inspection_id, auth_headers):
        r = requests.get(f"{API}/inspections", headers=auth_headers, timeout=15)
        assert r.status_code == 200
        ids = [d["id"] for d in r.json()]
        assert inspection_id in ids

    def test_full_doc_has_photos(self, inspection_id, auth_headers):
        r = requests.get(f"{API}/inspections/{inspection_id}", headers=auth_headers, timeout=15)
        d = r.json()
        assert len(d["photos"]) >= 10
        p0 = d["photos"][0]
        assert "id" in p0 and "original_filename" in p0

    def test_download_xlsx_valid(self, inspection_id, token, auth_headers):
        r = requests.get(f"{API}/inspections/{inspection_id}/download", params={"auth": token}, timeout=60)
        assert r.status_code == 200
        assert "openxml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        # one sheet per product
        full = requests.get(f"{API}/inspections/{inspection_id}", headers=auth_headers, timeout=15).json()
        assert len(wb.sheetnames) == len(full["products"])
        ws = wb[wb.sheetnames[0]]
        # merged cells preserved (~315)
        assert len(ws.merged_cells.ranges) > 250, f"merges={len(ws.merged_cells.ranges)}"
        # formula preserved
        assert str(ws["B30"].value).startswith("="), f"B30={ws['B30'].value}"
        # scale
        assert ws.page_setup.scale == 62 or str(ws.page_setup.scale) == "62"
        # embedded images in INSPECTION DIGITALS
        assert len(ws._images) > 0
        # old example data gone from A88
        a88 = (ws["A88"].value or "")
        assert "Open stitch" not in str(a88)
        # D210 is Style# ... for this run
        d210 = str(ws["D210"].value or "")
        assert d210.startswith("Style#") or "Style" in d210

    def test_photo_endpoint(self, inspection_id, token, auth_headers):
        full = requests.get(f"{API}/inspections/{inspection_id}", headers=auth_headers, timeout=15).json()
        pid = full["photos"][0]["id"]
        r = requests.get(f"{API}/inspections/{inspection_id}/photos/{pid}", params={"auth": token}, timeout=30)
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("image/")
        assert len(r.content) > 100
